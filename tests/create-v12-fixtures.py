"""Synthetic-only documents for Mac WPS v1.2 acceptance (python-docx)."""
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent

def save(doc, name):
    doc.save(OUT / ('WPS排版v12' + name + '.docx'))

def numbered(doc, text, num_id):
    p = doc.add_paragraph(text)
    pr = p._p.get_or_add_pPr()
    num = OxmlElement('w:numPr')
    level = OxmlElement('w:ilvl'); level.set(qn('w:val'), '0'); num.append(level)
    ident = OxmlElement('w:numId'); ident.set(qn('w:val'), str(num_id)); num.append(ident)
    pr.append(num)
    ind = OxmlElement('w:ind'); ind.set(qn('w:left'), '720'); ind.set(qn('w:hanging'), '360'); pr.append(ind)
    return p

def list_id(doc, start=1):
    root = doc.part.numbering_part.element
    abstract = OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'), '40')
    level = OxmlElement('w:lvl'); level.set(qn('w:ilvl'), '0')
    for tag, value in [('start', '1'), ('numFmt', 'decimal'), ('lvlText', '%1.'), ('lvlJc', 'left')]:
        node = OxmlElement('w:' + tag); node.set(qn('w:val'), value); level.append(node)
    abstract.append(level)
    if not root.xpath('./w:abstractNum[@w:abstractNumId="40"]'): root.append(abstract)
    number = OxmlElement('w:num'); number.set(qn('w:numId'), str(40 + start))
    aid = OxmlElement('w:abstractNumId'); aid.set(qn('w:val'), '40'); number.append(aid)
    override = OxmlElement('w:lvlOverride'); override.set(qn('w:ilvl'), '0')
    first = OxmlElement('w:startOverride'); first.set(qn('w:val'), str(start)); override.append(first)
    number.append(override); root.append(number)
    return 40 + start

def textbox(doc):
    # A native text box story verifies that body enumeration excludes its text.
    from lxml import etree
    p = doc.add_paragraph('文本框锚点保持内容')
    xml = '''<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape id="synthetic-box" style="width:120pt;height:35pt" type="#_x0000_t202"><v:textbox><w:txbxContent><w:p><w:r><w:t>合成文本框保持原样</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>'''
    p._p.append(etree.fromstring(xml))

notice = Document()
notice.add_paragraph('合成通知排版验收')
notice.add_paragraph('')
notice.add_paragraph('测试收件单位：')
notice.add_paragraph('通知正文第一段，部分文字选区按整个段落处理。')
notice.add_paragraph('此段保持原样，字体、缩进、换行和分页全部排除。')
notice.add_paragraph('第二个冒号段：')
notice.add_paragraph('后续正文继续使用正文缩进。')
notice.add_paragraph('')
notice.add_paragraph('合成发文单位')
notice.add_paragraph('')
notice.add_paragraph('2026年10月3日')
save(notice, '通知')

contract = Document()
contract.add_paragraph('合成合同验收', 'Title')
contract.add_paragraph('合同导语，样式识别回退为正文。')
contract.add_heading('第一章 总则', 1)
contract.add_paragraph('第一章正文。')
contract.add_heading('第一条 约定', 2)
for level in range(3, 10): contract.add_heading('第%d级标题' % level, level)
num = list_id(contract)
numbered(contract, '列表第一项，悬挂缩进保留。', num)
numbered(contract, '列表第二项。', num)
numbered(contract, '重新从五开始的列表项。', list_id(contract, 5))
table = contract.add_table(rows=2, cols=2)
for row in table.rows:
    for cell in row.cells: cell.text = '合成表格保持原样'
textbox(contract)
contract.add_paragraph('合成甲方')
contract.add_paragraph('合成乙方')
save(contract, '合同与列表')

attachments = Document()
attachments.add_paragraph('带附件文书')
attachments.add_heading('一、正文', 1)
attachments.add_paragraph('正文内容。')
attachments.add_paragraph('附件：')
attachments.add_paragraph('1. 合成附件说明')
attachments.add_page_break()
attachments.add_heading('附件标题', 1)
attachments.add_paragraph('较长段落。' * 600)
attachments.add_paragraph('合成单位')
attachments.add_paragraph('2026年10月3日')
save(attachments, '附件与长段')

sections = Document()
sections.add_paragraph('多节文档验收')
sections.add_heading('第一节标题', 1)
sections.add_paragraph('第一节正文。')
sections.add_section(WD_SECTION.NEW_PAGE)
sections.add_heading('第二节标题', 1)
sections.add_paragraph('第二节正文。')
sections.sections[0].header.paragraphs[0].text = '合成原页眉'
sections.sections[0].footer.paragraphs[0].text = '合成原页脚'
sections.add_paragraph('合成单位')
sections.add_paragraph('2026年10月3日')
save(sections, '多节')
print('已生成四份合成验收文档。')
