from docx import Document
from docx.enum.section import WD_SECTION
from pathlib import Path
out=Path.cwd()/'tests'
for name,multi in [('页眉页脚单节',False),('页眉页脚多节',True)]:
 d=Document();d.add_paragraph('排版功能验收');d.add_paragraph('仅用于插件测试，不包含真实案件内容。')
 d.add_page_break();d.add_paragraph('第二页正文')
 if multi:d.add_section(WD_SECTION.NEW_PAGE)
 else:d.add_page_break()
 d.add_paragraph('第三页正文');d.add_paragraph('测试单位');d.add_paragraph('2026年9月28日')
 d.save(out/('WPS排版'+name+'.docx'))

for name,complex_footer in [('单页替换',False),('复杂页脚保护',True)]:
 d=Document();d.add_paragraph('单页文书验收');d.add_paragraph('正文内容保持不变');d.add_paragraph('测试单位');d.add_paragraph('2026年9月28日')
 d.sections[0].header.paragraphs[0].text='原有页眉';d.sections[0].footer.paragraphs[0].text='原有页脚'
 if complex_footer:d.sections[0].footer.add_table(rows=1,cols=1,width=d.sections[0].page_width).cell(0,0).text='不能自动替换的表格'
 d.save(out/('WPS排版'+name+'.docx'))
